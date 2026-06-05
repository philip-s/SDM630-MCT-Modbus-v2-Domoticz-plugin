#!/usr/bin/env python3
"""insert_into_docx.py — (opcjonalnie) wstaw podgląd do strony tytułowej .docx.

Strona tytułowa jest w Wordzie, więc domyślnie podgląd wstawiasz ręcznie.
Ten skrypt automatyzuje to, jeśli w szablonie umieścisz placeholder tekstowy,
np. {{PODGLAD}} w akapicie, gdzie ma trafić obraz:

  insert_into_docx.py --docx strona_tytulowa.docx --image previews/dom.jpg \
                      --placeholder "{{PODGLAD}}" --width-cm 12 -o strona_gotowa.docx

Tryb alternatywny --replace-first podmienia PIERWSZY istniejący obraz w dokumencie
(wygodne, gdy w szablonie jest już obrazek-zaślepka):

  insert_into_docx.py --docx szablon.docx --image previews/dom.jpg --replace-first -o out.docx

Wymaga python-docx (pip install python-docx).
"""
from __future__ import annotations

import argparse
import os
import subprocess
import sys


def _ensure_docx():
    try:
        import docx  # noqa: F401
        return
    except Exception:
        pass
    try:
        subprocess.run([sys.executable, "-m", "pip", "install", "--quiet", "python-docx"], check=True)
    except Exception:
        sys.stderr.write("Brak python-docx. Zainstaluj: pip install python-docx\n")
        sys.exit(2)


_ensure_docx()

from docx import Document  # noqa: E402
from docx.shared import Cm  # noqa: E402


def insert_at_placeholder(doc, placeholder: str, image: str, width_cm: float | None) -> bool:
    for para in doc.paragraphs:
        if placeholder in para.text:
            # Wyczyść akapit i wstaw obraz w jednym runie.
            for run in list(para.runs):
                run.text = ""
            run = para.add_run()
            if width_cm:
                run.add_picture(image, width=Cm(width_cm))
            else:
                run.add_picture(image)
            return True
    return False


def replace_first_image(doc, image: str) -> bool:
    """Podmień dane bajtowe pierwszego osadzonego obrazu (zachowuje rozmiar/pozycję)."""
    part = doc.part
    with open(image, "rb") as fh:
        blob = fh.read()
    for rel_id, rel in part.rels.items():
        if "image" in rel.reltype:
            rel.target_part._blob = blob
            return True
    return False


def main(argv=None):
    ap = argparse.ArgumentParser(description="Wstaw podgląd JPG do strony tytułowej .docx.")
    ap.add_argument("--docx", required=True, help="Wejściowy plik .docx (szablon strony tytułowej).")
    ap.add_argument("--image", required=True, help="Plik podglądu (JPG/PNG) do wstawienia.")
    ap.add_argument("--output", "-o", required=True, help="Wynikowy plik .docx.")
    ap.add_argument("--placeholder", default="{{PODGLAD}}", help="Tekst-placeholder do podmiany. Domyślnie {{PODGLAD}}.")
    ap.add_argument("--width-cm", type=float, default=None, help="Szerokość wstawianego obrazu w cm (placeholder).")
    ap.add_argument("--replace-first", action="store_true", help="Zamiast placeholdera podmień pierwszy istniejący obraz.")
    args = ap.parse_args(argv)

    doc = Document(args.docx)
    if args.replace_first:
        ok = replace_first_image(doc, args.image)
        how = "podmieniono pierwszy obraz"
    else:
        ok = insert_at_placeholder(doc, args.placeholder, args.image, args.width_cm)
        how = f"wstawiono w miejsce '{args.placeholder}'"
    if not ok:
        sys.stderr.write(
            "Nie znaleziono celu wstawienia. "
            "Dodaj placeholder w szablonie albo użyj --replace-first.\n"
        )
        sys.exit(1)

    os.makedirs(os.path.dirname(os.path.abspath(args.output)) or ".", exist_ok=True)
    doc.save(args.output)
    print(f"OK: {args.output}  ({how})")


if __name__ == "__main__":
    main()
